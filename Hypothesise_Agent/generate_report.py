#!/usr/bin/env python3
"""Generates a Word document project report for the Hypothesise Agent project."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime
import os

OUTPUT_PATH = os.path.join(os.path.dirname(__file__), "Hypothesise_Agent_Project_Report.docx")

# ── Helpers ──────────────────────────────────────────────────────────────────

def set_cell_bg(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)

def add_heading(doc, text, level=1, color=None):
    h = doc.add_heading(text, level=level)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if color:
        for run in h.runs:
            run.font.color.rgb = RGBColor(*bytes.fromhex(color))
    return h

def add_para(doc, text, bold=False, italic=False, size=11, color=None, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(1)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*bytes.fromhex(color))
    return p

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.left_indent = Cm(1 + level * 0.5)
    return p

def add_numbered(doc, text):
    return doc.add_paragraph(text, style="List Number")

def add_divider(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CCCCCC")
    pBdr.append(bottom)
    pPr.append(pBdr)

def styled_table(doc, headers, rows, header_bg="1E3A5F", header_fg="FFFFFF", alt_bg="EAF1FB"):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header row
    hdr_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.text = h
        set_cell_bg(cell, header_bg)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(*bytes.fromhex(header_fg))
            run.font.size = Pt(10)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Data rows
    for ri, row in enumerate(rows):
        tr = table.rows[ri + 1]
        bg = alt_bg if ri % 2 == 0 else "FFFFFF"
        for ci, val in enumerate(row):
            cell = tr.cells[ci]
            cell.text = str(val)
            set_cell_bg(cell, bg)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(10)
    return table

# ── Document ─────────────────────────────────────────────────────────────────

doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(2.5)

# ── Cover ──────────────────────────────────────────────────────────────────────

doc.add_paragraph()
doc.add_paragraph()

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("HYPOTHESISE AGENT")
run.bold = True
run.font.size = Pt(28)
run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_run = subtitle.add_run("AI-Powered Scientific Research & Experimentation Platform")
sub_run.font.size = Pt(14)
sub_run.italic = True
sub_run.font.color.rgb = RGBColor(0x44, 0x72, 0xC4)

doc.add_paragraph()

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run(f"Project Report  ·  {datetime.date.today().strftime('%B %d, %Y')}").font.size = Pt(11)

doc.add_paragraph()
add_divider(doc)
doc.add_paragraph()

# ── 1. Executive Summary ───────────────────────────────────────────────────────

add_heading(doc, "1. Executive Summary", level=1, color="1E3A5F")
add_para(doc,
    "Hypothesise Agent is a multi-agent AI platform that automates the full scientific "
    "research lifecycle — from hypothesis generation and literature review, through "
    "experimental design, to simulation, numerical analysis, and iterative learning. "
    "Researchers describe any experiment in plain language; the system autonomously "
    "searches academic papers, produces a rigorous experimental design, runs a domain-"
    "specific simulation, generates publication-quality plots, and stores feedback-"
    "weighted memories for continuous improvement.")

doc.add_paragraph()

# ── 2. Project Overview ────────────────────────────────────────────────────────

add_heading(doc, "2. Project Overview", level=1, color="1E3A5F")

add_heading(doc, "2.1  Purpose & Vision", level=2)
add_para(doc,
    "Scientific research traditionally requires weeks of literature review, expert "
    "knowledge to design experiments, and expensive lab resources for simulation. "
    "Hypothesise Agent condenses this into seconds by combining real-time paper "
    "retrieval (Bohrium Science Navigator + arXiv), large-language-model reasoning "
    "(Ollama / DeepSeek / Claude), and deterministic numerical simulation engines "
    "across seven scientific domains.")

add_heading(doc, "2.2  Key Capabilities", level=2)
for item in [
    "Natural-language experiment description → fully structured hypothesis",
    "Live academic paper search from Bohrium AI and arXiv",
    "Automated experimental design with variables, controls, and procedure steps",
    "Domain-aware simulation (chemistry, biology, physics, civil, mechanical, electronics, CS, general)",
    "Real-time streaming responses via Server-Sent Events (SSE)",
    "Interactive React dashboard with Plotly charts",
    "Memory & feedback loop — scores are stored and used to reweight future retrieval",
    "REST API with auto-generated OpenAPI/Swagger docs",
]:
    add_bullet(doc, item)

doc.add_paragraph()
add_divider(doc)

# ── 3. System Architecture ─────────────────────────────────────────────────────

add_heading(doc, "3. System Architecture", level=1, color="1E3A5F")
add_para(doc,
    "The platform follows a layered multi-agent pipeline. Each layer is an independent "
    "service or module that can be replaced or extended without affecting the others.")

doc.add_paragraph()

arch_headers = ["Layer", "Component", "Technology", "Port"]
arch_rows = [
    ["L1", "Hypothesis Agent",         "Python · LiteLLM · DeepSeek",         "CLI / API"],
    ["L2", "Experimental Design Agent","Python · Ollama (llama3)",             "CLI / API"],
    ["L3", "Simulation Agent",         "Python · NumPy · SciPy",              "Internal"],
    ["L4", "Backend API",              "FastAPI · Uvicorn · SSE",             "8000"],
    ["L5", "Frontend UI",              "React 18 · Vite · Plotly.js",         "3000"],
    ["L6", "Memory Agent",             "FastAPI · SQLite · ChromaDB (stub)",   "8001"],
    ["L7", "Vector Store",             "ChromaDB local / Pinecone cloud",      "N/A"],
]
styled_table(doc, arch_headers, arch_rows)

doc.add_paragraph()
add_divider(doc)

# ── 4. Agent Descriptions ──────────────────────────────────────────────────────

add_heading(doc, "4. Agent Descriptions", level=1, color="1E3A5F")

# 4.1
add_heading(doc, "4.1  Hypothesis Agent  (hypo_agent.py)", level=2)
add_para(doc,
    "The entry point for generating structured scientific hypotheses. It accepts any "
    "plain-language experiment description, searches Bohrium Science Navigator and "
    "arXiv for related papers, and passes them to DeepSeek LLM to produce a structured "
    "hypothesis with background, variables, methodology, and expected outcomes. "
    "Results are saved as JSON and a human-readable report under output/.")
add_bullet(doc, "LLM: DeepSeek via LiteLLM (configurable via MODEL env var)")
add_bullet(doc, "Parameters: temperature=0.4, max_tokens=4096, top_p=0.9, frequency_penalty=0.1")
add_bullet(doc, "Dry-run mode available for offline testing")

doc.add_paragraph()

# 4.2
add_heading(doc, "4.2  Experimental Design Agent  (experimental_design_agent.py)", level=2)
add_para(doc,
    "Takes a research topic and produces a full experimental protocol. It queries "
    "Bohrium for domain-specific literature, analyses methodology gaps, then uses "
    "Ollama (llama3 by default) to draft a detailed design covering independent / "
    "dependent variables, control groups, measurement methods, and expected limitations.")
add_bullet(doc, "LLM: Ollama local (llama3 default, configurable via OLLAMA_MODEL)")
add_bullet(doc, "Parameters: temperature=0.3, top_p=0.9, num_predict=4096, repeat_penalty=1.1")
add_bullet(doc, "Timeout: 180 s to handle complex multi-step prompts")

doc.add_paragraph()

# 4.3
add_heading(doc, "4.3  Simulation Agent  (simulation_agent/)", level=2)
add_para(doc,
    "Receives an experimental procedure JSON and routes it to the most appropriate "
    "deterministic engine. If no domain engine matches, it falls back to the LLM "
    "simulation engine (Claude or Ollama). Returns structured results with numerical "
    "values, plot data, observations, a confidence score, and a step-by-step log.")

doc.add_paragraph()
sim_headers = ["Engine", "Domain", "Key Libraries"]
sim_rows = [
    ["chem_sim",        "Chemistry / Molecular",     "NumPy · SciPy · RDKit (optional)"],
    ["bio_sim",         "Biology / Genomics",         "NumPy · SciPy · BioPython (optional)"],
    ["physics_sim",     "Physics",                    "NumPy · SciPy"],
    ["civil_sim",       "Civil / Structural Eng.",    "NumPy · SciPy"],
    ["mech_sim",        "Mechanical Engineering",     "NumPy · SciPy"],
    ["electronics_sim", "Electronics / Arduino",      "NumPy"],
    ["cs_sim",          "Computer Science / ML",      "NumPy"],
    ["llm_fallback",    "General / Any domain",       "Claude API · Ollama"],
]
styled_table(doc, sim_headers, sim_rows)
doc.add_paragraph()

# 4.4
add_heading(doc, "4.4  Backend API  (backend/server.py)", level=2)
add_para(doc,
    "A FastAPI service that exposes REST endpoints consumed by the React frontend. "
    "Experiment design is streamed in real-time via Server-Sent Events so the user "
    "sees tokens appearing as they are generated. After streaming completes, the "
    "simulation agent is invoked synchronously and the full result is returned as JSON.")
add_bullet(doc, "POST /api/design-stream  — SSE stream of experimental design text")
add_bullet(doc, "POST /api/simulate       — Run domain simulation, return JSON results")
add_bullet(doc, "GET  /api/health         — Health check")
add_bullet(doc, "GET  /docs               — Swagger UI (auto-generated)")
add_bullet(doc, "CORS enabled for localhost:3000, 5173")

doc.add_paragraph()

# 4.5
add_heading(doc, "4.5  Frontend UI  (frontend/)", level=2)
add_para(doc,
    "A React 18 single-page application built with Vite. The user types an experiment "
    "topic, selects optional domain hints, and submits. The design streams token-by-"
    "token into a markdown renderer. Once complete, simulation results appear in a "
    "structured panel containing numerical results, interactive Plotly charts, an "
    "observation section, and a collapsible simulation log.")
add_bullet(doc, "Observation section: unique amber/gold gradient background with hover slide effect")
add_bullet(doc, "Charts: react-plotly.js — line, scatter, bar, histogram")
add_bullet(doc, "Expandable full-screen plot view")
add_bullet(doc, "Confidence score badge and domain badge per simulation run")

doc.add_paragraph()

# 4.6
add_heading(doc, "4.6  Memory Agent  (memory_agent/)", level=2)
add_para(doc,
    "An independent FastAPI microservice running on port 8001. It records quality "
    "feedback scores for every recalled memory item, tracks per-layer pipeline "
    "accuracy over time, dynamically reweights retrieval based on performance, and "
    "generates improvement plans using Claude LangChain integration.")

doc.add_paragraph()
mem_headers = ["Store", "Technology", "Tables / Collections", "Status"]
mem_rows = [
    ["SQLite DB",        "sqlite3 (stdlib)",   "memory_feedback, run_accuracy, retrieval_weights", "Active"],
    ["Vector Store",     "ChromaDB (stub)",    "chat_history, experiment_results, hypothesis_procedures, agent_conversations, model_learnings", "Stub — ready to wire"],
    ["Cloud Vector",     "Pinecone (optional)","Same collections via cloud_available() flag",       "Disabled"],
]
styled_table(doc, mem_headers, mem_rows)

doc.add_paragraph()
add_divider(doc)

# ── 5. Technology Stack ────────────────────────────────────────────────────────

add_heading(doc, "5. Technology Stack", level=1, color="1E3A5F")

doc.add_paragraph()
tech_headers = ["Category", "Technology", "Version / Notes"]
tech_rows = [
    ["Language",         "Python",             "3.10+ (tested on 3.13)"],
    ["Language",         "JavaScript / JSX",   "ES2022"],
    ["Web Framework",    "FastAPI",             "0.133+"],
    ["ASGI Server",      "Uvicorn",             "0.41+"],
    ["Frontend",         "React",              "18.2"],
    ["Build Tool",       "Vite",               "5.0"],
    ["Charts",           "Plotly.js",          "2.29"],
    ["Markdown",         "react-markdown",     "9.0"],
    ["LLM (cloud)",      "DeepSeek Chat",      "via LiteLLM"],
    ["LLM (cloud)",      "Claude Sonnet 4",    "claude-sonnet-4-20250514"],
    ["LLM (local)",      "Ollama llama3",      "http://localhost:11434"],
    ["LLM orchestration","LiteLLM",            "any-provider abstraction"],
    ["LLM orchestration","LangChain Anthropic","memory improvement plans"],
    ["Numerics",         "NumPy",              "2.4+"],
    ["Numerics",         "SciPy",              "1.17+"],
    ["Database",         "SQLite",             "stdlib — memory_feedback.db"],
    ["Vector DB",        "ChromaDB",           "stub (PersistentClient ready)"],
    ["Papers",           "Bohrium Science Nav","REST API"],
    ["Papers",           "arXiv",              "arxiv Python client"],
    ["Chemistry",        "RDKit",              "optional"],
    ["Biology",          "BioPython",          "optional"],
    ["Env vars",         "python-dotenv",      ".env file"],
    ["Report gen",       "python-docx",        "this report"],
]
styled_table(doc, tech_headers, tech_rows)

doc.add_paragraph()
add_divider(doc)

# ── 6. Data Flow ───────────────────────────────────────────────────────────────

add_heading(doc, "6. End-to-End Data Flow", level=1, color="1E3A5F")
add_para(doc, "A typical user interaction follows these steps:")
doc.add_paragraph()

steps = [
    ("User Input",           "User types experiment topic in React UI (port 3000)"),
    ("API Request",          "Frontend sends POST /api/design-stream to FastAPI backend (port 8000)"),
    ("Paper Retrieval",      "Backend queries Bohrium Science Navigator + arXiv for related papers"),
    ("LLM Design",           "Ollama (llama3) generates experimental design; tokens stream back via SSE"),
    ("UI Streaming",         "React renders design markdown token-by-token in real time"),
    ("Simulation Trigger",   "Frontend sends POST /api/simulate with the completed design text"),
    ("Domain Routing",       "Simulation Agent parses design, detects domain, routes to correct engine"),
    ("Numerical Simulation", "Engine (e.g. physics_sim) runs SciPy/NumPy calculations → returns results + plots"),
    ("LLM Fallback",         "If engine fails or domain = general, Claude API or Ollama generates results"),
    ("Results Display",      "Frontend renders numerical results, Plotly charts, observations, confidence score"),
    ("Memory Storage",       "Feedback & accuracy scores persisted to SQLite memory_feedback.db"),
    ("Retrieval Reweighting","Memory agent updates retrieval_weights table for future improvement"),
]

flow_headers = ["#", "Stage", "Description"]
flow_rows = [(str(i+1), stage, desc) for i, (stage, desc) in enumerate(steps)]
styled_table(doc, flow_headers, flow_rows)

doc.add_paragraph()
add_divider(doc)

# ── 7. Project Structure ───────────────────────────────────────────────────────

add_heading(doc, "7. Project File Structure", level=1, color="1E3A5F")
doc.add_paragraph()

structure = """\
hypothesis_generator/
├── hypo_agent.py                  ← Hypothesis generation agent (DeepSeek/LiteLLM)
├── experimental_design_agent.py   ← Experimental design via Ollama
├── launch.py / launch.ps1         ← One-command project launcher
├── .env                           ← API keys & config (not committed)
│
├── backend/
│   ├── server.py                  ← FastAPI API, SSE streaming, simulation trigger
│   └── requirements.txt
│
├── frontend/
│   ├── src/
│   │   ├── App.jsx                ← Main React app, state management
│   │   ├── components/
│   │   │   ├── SimulationResults.jsx   ← Results panel (charts, obs, scores)
│   │   │   └── SimulationResults.css   ← Unique amber observation styling
│   │   └── index.css
│   ├── package.json
│   └── vite.config.js
│
├── simulation_agent/
│   ├── agent.py                   ← Orchestrator: parse → route → simulate → format
│   ├── router.py                  ← Domain → engine mapping
│   ├── schema.py                  ← Pydantic models
│   ├── output_formatter.py
│   └── engines/
│       ├── chem_sim.py            ← Chemistry (RDKit optional)
│       ├── bio_sim.py             ← Biology (BioPython optional)
│       ├── physics_sim.py         ← Physics
│       ├── civil_sim.py           ← Structural/Civil engineering
│       ├── mech_sim.py            ← Mechanical engineering
│       ├── electronics_sim.py     ← Electronics/Arduino
│       ├── cs_sim.py              ← CS/ML experiments
│       └── llm_fallback.py        ← Claude/Ollama fallback
│
├── memory_agent/
│   ├── main.py                    ← FastAPI service (port 8001)
│   ├── memory_feedback.db         ← SQLite database (local)
│   ├── core/
│   │   ├── feedback.py            ← SQLite store: scores, accuracy, weights
│   │   ├── improvement_engine.py  ← LangChain/Claude improvement plans
│   │   ├── chromadb_store.py      ← Vector store stub (ready to activate)
│   │   ├── embedder.py            ← Embedding stub
│   │   └── accuracy_tracker.py
│   └── routers/
│       └── improve.py             ← /improve endpoints
│
└── output/                        ← Saved experiment designs (timestamped .txt)"""

p = doc.add_paragraph()
run = p.add_run(structure)
run.font.name = "Courier New"
run.font.size = Pt(8.5)

doc.add_paragraph()
add_divider(doc)

# ── 8. Configuration ───────────────────────────────────────────────────────────

add_heading(doc, "8. Configuration & Environment Variables", level=1, color="1E3A5F")
doc.add_paragraph()

cfg_headers = ["Variable", "Default", "Description"]
cfg_rows = [
    ["OLLAMA_URL",        "http://localhost:11434/api/generate", "Ollama API endpoint"],
    ["OLLAMA_MODEL",      "llama3",                             "Model to use for Ollama"],
    ["OLLAMA_TEMPERATURE","0.3",                                "Generation temperature"],
    ["DEEPSEEK_API_KEY",  "(none)",                             "DeepSeek API key for hypo_agent"],
    ["MODEL",             "deepseek/deepseek-chat",             "LiteLLM model string"],
    ["ANTHROPIC_API_KEY", "(none)",                             "Claude API key for fallback + memory"],
]
styled_table(doc, cfg_headers, cfg_rows)
add_para(doc, "Create a .env file in the hypothesis_generator/ directory to set these values.", italic=True)

doc.add_paragraph()
add_divider(doc)

# ── 9. How to Run ──────────────────────────────────────────────────────────────

add_heading(doc, "9. How to Run the Project", level=1, color="1E3A5F")

add_heading(doc, "9.1  Prerequisites", level=2)
add_bullet(doc, "Python 3.10+ with virtual environment activated")
add_bullet(doc, "Node.js 18+ and npm")
add_bullet(doc, "Ollama installed and running:  ollama serve  +  ollama pull llama3")
add_bullet(doc, "Optional: DEEPSEEK_API_KEY and ANTHROPIC_API_KEY in .env")

add_heading(doc, "9.2  Quick Start (Recommended)", level=2)
p = doc.add_paragraph()
r = p.add_run("    .\\launch.ps1           # Windows PowerShell\n    python launch.py        # cross-platform")
r.font.name = "Courier New"
r.font.size = Pt(10)

add_heading(doc, "9.3  Manual Start", level=2)
steps_manual = [
    "# Terminal 1 — Backend",
    "cd backend",
    "python -m uvicorn server:app --port 8000",
    "",
    "# Terminal 2 — Frontend",
    "cd frontend",
    "npm run dev",
    "",
    "# Terminal 3 — Memory Agent (optional)",
    "cd memory_agent",
    "python main.py",
]
p = doc.add_paragraph()
r = p.add_run("\n".join(steps_manual))
r.font.name = "Courier New"
r.font.size = Pt(10)

add_heading(doc, "9.4  URLs", level=2)
url_headers = ["Service", "URL"]
url_rows = [
    ["Frontend UI",    "http://localhost:3000"],
    ["Backend API",    "http://localhost:8000"],
    ["API Docs",       "http://localhost:8000/docs"],
    ["Memory Agent",   "http://localhost:8001"],
]
styled_table(doc, url_headers, url_rows)

doc.add_paragraph()
add_divider(doc)

# ── 10. Model Optimization Summary ────────────────────────────────────────────

add_heading(doc, "10. LLM Model Optimization", level=1, color="1E3A5F")
add_para(doc, "All LLM calls have been optimized with the following parameters:")
doc.add_paragraph()

opt_headers = ["File", "Model", "max_tokens", "temperature", "top_p", "Extra"]
opt_rows = [
    ["experimental_design_agent.py", "Ollama llama3",          "4096", "0.3",  "0.9",  "repeat_penalty=1.1, num_predict=4096, timeout=180s"],
    ["simulation_agent/llm_fallback.py (Claude)", "claude-sonnet-4-20250514", "4096", "0.3", "0.9", ""],
    ["simulation_agent/llm_fallback.py (Ollama)", "Ollama llama3", "4096 (num_predict)", "0.3", "0.9", "timeout=180s"],
    ["hypo_agent.py",                "deepseek-chat",           "4096", "0.4",  "0.9",  "frequency_penalty=0.1"],
    ["memory_agent/improvement_engine.py", "claude-sonnet-4-20250514", "2048", "0",  "0.95", ""],
]
styled_table(doc, opt_headers, opt_rows)

doc.add_paragraph()
add_divider(doc)

# ── 11. Future Enhancements ────────────────────────────────────────────────────

add_heading(doc, "11. Future Enhancements", level=1, color="1E3A5F")

enhancements = [
    ("Real Vector Store",      "Activate ChromaDB PersistentClient or connect Pinecone for semantic memory recall"),
    ("Real Embedder",          "Replace stub embedder with sentence-transformers or OpenAI embeddings"),
    ("RDKit & BioPython",       "Install optional packages to enable full chemistry and biology simulations"),
    ("Multi-model Support",    "Add GPT-4o support via LiteLLM for broader model choice"),
    ("Experiment History",     "UI page showing past experiments with replay and comparison"),
    ("Export Functionality",   "Export simulation results as PDF or CSV directly from the UI"),
    ("Collaborative Mode",     "Multi-user sessions with shared experiment workspace"),
    ("Cloud Deployment",       "Dockerise all services; deploy to AWS/GCP with Nginx reverse proxy"),
    ("Unit Tests",             "Expand test coverage in tests/ for all simulation engines"),
]

for title_e, desc in enhancements:
    p = doc.add_paragraph(style="List Bullet")
    run_t = p.add_run(f"{title_e}: ")
    run_t.bold = True
    run_t.font.size = Pt(11)
    run_d = p.add_run(desc)
    run_d.font.size = Pt(11)

doc.add_paragraph()
add_divider(doc)

# ── 12. Summary ────────────────────────────────────────────────────────────────

add_heading(doc, "12. Summary", level=1, color="1E3A5F")
add_para(doc,
    "Hypothesise Agent demonstrates how modern AI tooling — large language models, "
    "domain-specific numerical engines, vector memory, and streaming web APIs — can "
    "be composed into a cohesive research assistant. The modular architecture means "
    "each layer (LLM, simulation, memory, UI) can be upgraded independently. "
    "With minimal configuration (an Ollama instance and optional API keys), researchers "
    "across chemistry, biology, physics, engineering, and computer science can generate "
    "rigorous experimental designs and simulations in seconds rather than days.")

doc.add_paragraph()

footer_p = doc.add_paragraph()
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer_p.add_run(f"Generated on {datetime.date.today().strftime('%B %d, %Y')}  ·  Hypothesise Agent Project Report")
fr.font.size = Pt(9)
fr.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
fr.italic = True

# ── Save ───────────────────────────────────────────────────────────────────────

doc.save(OUTPUT_PATH)
print(f"✅  Report saved to: {OUTPUT_PATH}")
